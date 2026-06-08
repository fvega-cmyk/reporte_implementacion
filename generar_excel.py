"""
Genera el Excel del reporte con hojas 'Resumen' y 'BBDD'.
Equivalente a generarExcel() en Apps Script.
Usa openpyxl en vez de copiar el template de Google Sheets.
"""
from io import BytesIO
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import IDX, INDICES_FECHA, INDICES_FOTO, ESTADOS
from utils import fmt, foto_url


# Estilos reutilizables
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
TOTAL_FILL = PatternFill("solid", fgColor="1F3864")
TOTAL_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ALT_FILL = PatternFill("solid", fgColor="EEF2F9")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LINK_FONT = Font(name="Calibri", size=11, color="0563C1", underline="single")


def _aplicar_estilos_encabezado(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER


def _ajustar_anchos(ws, ncols, ancho_default=18):
    for c in range(1, ncols + 1):
        ws.column_dimensions[get_column_letter(c)].width = ancho_default


def construir_resumen(ws, filas):
    """Construye la hoja Resumen (espejo de la sección Resumen en GS)."""
    # Encabezado del bloque informativo
    ws["A1"] = "Reporte Diario de Implementación"
    ws["A1"].font = Font(size=14, bold=True, color="1F3864")
    ws.merge_cells("A1:I1")

    ws["A3"] = "Fecha inicio campaña:"
    ws["A4"] = "Fecha término campaña:"
    ws["A3"].font = Font(bold=True)
    ws["A4"].font = Font(bold=True)
    ws["B3"] = fmt(filas[0][IDX["FECHA_INICIO"]])
    ws["B4"] = fmt(filas[0][IDX["FECHA_TERMINO"]])

    # Encabezados de la tabla
    headers_tabla = ["Código", "Nombre Sala", "Comuna", "Fecha Entrega"] + ESTADOS + ["Total"]
    ncols = len(headers_tabla)
    for c, h in enumerate(headers_tabla, start=1):
        ws.cell(row=6, column=c, value=h)
    _aplicar_estilos_encabezado(ws, ncols, row=6)

    # Agrupar por sala
    salas = {}
    for row in filas:
        cod = (row[IDX["COD"]] or "").strip()
        if cod not in salas:
            salas[cod] = {
                "cod": cod,
                "nombre": row[IDX["NOMBRE_SALA"]] or "",
                "comuna": row[IDX["COMUNA"]] or "",
                "fecha_entrega": "",
                "conteos": defaultdict(int),
            }
        if not salas[cod]["fecha_entrega"] and row[IDX["FECHA_ENTREGA"]]:
            salas[cod]["fecha_entrega"] = fmt(row[IDX["FECHA_ENTREGA"]])
        proc = (row[IDX["PROCESO"]] or "").strip()
        if "reagenda" in proc.lower():
            proc = "Reagenda"
        salas[cod]["conteos"][proc] += 1

    # Volcar filas
    totales = {e: 0 for e in ESTADOS}
    total_general = 0
    fila_r = 7
    for cod in sorted(salas.keys()):
        s = salas[cod]
        ws.cell(row=fila_r, column=1, value=s["cod"]).alignment = CENTER
        ws.cell(row=fila_r, column=2, value=s["nombre"]).alignment = LEFT
        ws.cell(row=fila_r, column=3, value=s["comuna"]).alignment = LEFT
        ws.cell(row=fila_r, column=4, value=s["fecha_entrega"]).alignment = CENTER
        total_sala = 0
        for i, est in enumerate(ESTADOS, start=5):
            v = s["conteos"].get(est, 0)
            ws.cell(row=fila_r, column=i, value=(v if v > 0 else "")).alignment = CENTER
            totales[est] += v
            total_sala += v
        ws.cell(row=fila_r, column=4 + len(ESTADOS) + 1, value=total_sala).alignment = CENTER
        total_general += total_sala

        # Zebra striping
        if fila_r % 2 == 0:
            for c in range(1, ncols + 1):
                ws.cell(row=fila_r, column=c).fill = ALT_FILL
        for c in range(1, ncols + 1):
            ws.cell(row=fila_r, column=c).border = BORDER
        fila_r += 1

    # Fila Total general
    ws.cell(row=fila_r, column=1, value="Total general")
    for i, est in enumerate(ESTADOS, start=5):
        v = totales.get(est, 0)
        ws.cell(row=fila_r, column=i, value=(v if v > 0 else ""))
    ws.cell(row=fila_r, column=4 + len(ESTADOS) + 1, value=total_general)
    for c in range(1, ncols + 1):
        cell = ws.cell(row=fila_r, column=c)
        cell.fill = TOTAL_FILL
        cell.font = TOTAL_FONT
        cell.alignment = CENTER if c > 1 else LEFT
        cell.border = BORDER

    # Anchos
    anchos = [12, 32, 18, 14] + [14] * len(ESTADOS) + [10]
    for c, w in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[6].height = 30
    ws.freeze_panes = "A7"


def construir_bbdd(ws, headers, filas):
    """Construye la hoja BBDD: todas las columnas, fechas formateadas, fotos como hipervínculos."""
    ws.append(headers)
    _aplicar_estilos_encabezado(ws, len(headers), row=1)

    for r_off, row in enumerate(filas, start=2):
        for c, val in enumerate(row, start=0):
            idx = c
            if idx in INDICES_FECHA:
                valor = fmt(val)
            elif idx in INDICES_FOTO:
                # Hipervínculo "Ver FOTO N"
                ruta = (val or "").strip()
                if ruta:
                    url = foto_url(ruta)
                    num = idx - IDX["FOTO1"] + 1
                    cell = ws.cell(row=r_off, column=c + 1, value=f"Ver FOTO {num}")
                    cell.hyperlink = url
                    cell.font = LINK_FONT
                    cell.alignment = CENTER
                    continue
                else:
                    valor = ""
            else:
                valor = val
            cell = ws.cell(row=r_off, column=c + 1, value=valor)
            if idx in INDICES_FECHA:
                cell.alignment = CENTER

        if r_off % 2 == 0:
            for c in range(1, len(headers) + 1):
                if ws.cell(row=r_off, column=c).fill.fgColor.rgb in (None, "00000000"):
                    ws.cell(row=r_off, column=c).fill = ALT_FILL

    _ajustar_anchos(ws, len(headers), 18)
    ws.freeze_panes = "A2"
    # Habilitar filtros
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(filas) + 1}"


def generar_excel(campana, filas, headers, hoy):
    """
    Construye el Excel completo y devuelve los bytes.
    Equivalente a generarExcel() en Apps Script.
    """
    wb = Workbook()
    ws_resumen = wb.active
    ws_resumen.title = "Resumen"
    construir_resumen(ws_resumen, filas)

    ws_bbdd = wb.create_sheet("BBDD")
    construir_bbdd(ws_bbdd, headers, filas)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
